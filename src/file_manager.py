"""Report file management and export module.

This module handles the storage and export of generated sales reports and
structured analysis results.

It creates destination directories when necessary, generates a shared base
filename using the source CSV filename and a timestamp, saves human-readable
reports as TXT files, exports structured analysis results as JSON, saves
aggregated analysis summaries as independent CSV files, and generates
multi-sheet Excel workbooks.

Exported analysis data may include product, category, monthly, city, and
payment-method summaries, monthly growth metrics, monthly best-selling
products, monthly highest-income categories, Top 5 product rankings,
validation errors, and validation warnings.

JSON-compatible analysis structures convert pandas missing values into
`None` before serialization.

TXT and JSON files are written using UTF-8 encoding. CSV files are generated
from pandas DataFrames, while XLSX worksheets are populated from analysis
results using openpyxl.
"""


from pathlib import Path
from datetime import datetime
import numpy as np
import pandas as pd
import json
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.utils.dataframe import dataframe_to_rows

from typing import Dict, Any, List

from src.errors import ReportSaveError


def save_report(
        report_text: str,
        output_folder: str | Path,
        file_name: str
) -> Path:
    """Save a generated sales report as a TXT file.

    Creates the destination directory and any missing parent directories before
    writing the report.

    The function receives a previously generated shared base filename, adds the
    `.txt` extension, and writes the report content using UTF-8 encoding.

    Args:
        report_text: Complete human-readable report content to save.
        output_folder: Directory where the TXT report will be stored.
        file_name: Shared base filename without a file extension.

    Returns:
        Path to the generated TXT report.

    Raises:
        ReportSaveError: If the destination directory cannot be created or the
            report file cannot be written because of a file-system error.
    """
    output_filename = f"{file_name}.txt"
    try:
        folder = Path(output_folder)
        folder.mkdir(parents=True, exist_ok=True)
        path = folder / output_filename
        with open(path, "w", encoding="utf-8") as file:
            file.write(report_text)
    except OSError as error:
        raise ReportSaveError() from error
    return path

def create_report_base_name(input_file_path: str | Path) -> str:
    """Generate a shared base filename for generated output files.

    Extracts the source CSV filename without its extension and combines it
    with the current local date and time.

    The timestamp contains the year, month, day, hour, minute, second, and
    milliseconds.

    The resulting base filename is reused by report-file and chart-generation
    workflows so outputs created during the same execution share a common
    identifier.

    Args:
        input_file_path: Path of the source CSV file whose filename will be
            used as the output filename prefix.

    Returns:
        A base filename formatted as
        `<source_filename>_YYYY-MM-DD_HH-MM-SS-fff`.
    """
    file_path = Path(input_file_path).stem
    now = datetime.now()
    timestamp = now.strftime("%Y-%m-%d_%H-%M-%S-%f")[:-3]
    file_name = f"{file_path}_{timestamp}"
    return file_name

def save_analysis_json(
        analysis_result: Dict[str, Any], 
        output_folder: str | Path, 
        file_name: str
) -> Path:
    """Save the complete sales analysis result as a JSON file.

    Creates a shallow copy of the analysis result and converts pandas DataFrame
    structures into lists of dictionaries so they can be serialized to JSON.

    Before conversion, pandas `NaN` values are replaced with `None` so missing
    analysis values are represented as JSON `null` values.

    The following DataFrames are always converted:

    - `product_summary`
    - `category_summary`
    - `monthly_summary`
    - `monthly_best_selling_product`
    - `monthly_highest_income_category`

    Optional city and payment-method summaries are also converted when available.

    The function adds the `.json` extension to the shared base filename, creates
    the destination directory when necessary, and writes the resulting structure
    using UTF-8 encoding.

    Args:
        analysis_result: Dictionary containing the complete sales-analysis result.
        output_folder: Directory where the JSON analysis file will be stored.
        file_name: Shared base filename without a file extension.

    Returns:
        Path to the generated JSON analysis file.

    Raises:
        ReportSaveError: If the destination directory cannot be created or the
            JSON file cannot be written because of a file-system error.
    """
    analysis_json = analysis_result.copy()
    analysis_json["product_summary"] = analysis_json["product_summary"].replace({np.nan: None}).to_dict(orient="records")
    analysis_json["category_summary"] = analysis_json["category_summary"].replace({np.nan: None}).to_dict(orient="records")
    analysis_json["monthly_summary"] = analysis_json["monthly_summary"].replace({np.nan: None}).to_dict(orient="records")
    analysis_json["monthly_best_selling_product"] = analysis_json["monthly_best_selling_product"].replace({np.nan: None}).to_dict(orient="records")
    analysis_json["monthly_highest_income_category"] = analysis_json["monthly_highest_income_category"].replace({np.nan: None}).to_dict(orient="records")
    if "city_summary" in analysis_json.keys() and analysis_json["city_summary"] is not None:
        analysis_json["city_summary"] = analysis_json["city_summary"].replace({np.nan: None}).to_dict(orient="records")
    if "payment_method_summary" in analysis_json.keys() and analysis_json["payment_method_summary"] is not None:
        analysis_json["payment_method_summary"] = analysis_json["payment_method_summary"].replace({np.nan: None}).to_dict(orient="records")
    output_filename = f"{file_name}.json"
    try:
        folder = Path(output_folder)
        folder.mkdir(parents=True, exist_ok=True)
        path = folder / output_filename
        with open(path, "w", encoding="utf-8") as file:
            json.dump(analysis_json, file, indent=4, ensure_ascii=False)
    except OSError as error:
        raise ReportSaveError() from error
    return path

def save_analysis_result_csv_files(
        analysis_result: Dict[str, Any], 
        output_folder: str | Path, 
        file_name: str
) -> Dict[str, Path]:
    """Save analysis summary DataFrames as independent CSV files.

        Creates individual CSV files for product, category, monthly, monthly
        best-selling-product, and monthly highest-income-category analyses.

        Optional city and payment-method summaries are also exported when available.

        Each generated file uses the shared report base filename followed by a
        descriptive suffix identifying the corresponding analysis structure.

        The creation and storage of each individual CSV file is delegated to
        `create_save_analysis_result_csv_files_and_path()`.

        Args:
            analysis_result: Dictionary containing the complete sales-analysis result
                and its DataFrame structures.
            output_folder: Directory where the CSV files will be stored.
            file_name: Shared base filename without a file extension.

        Returns:
            A dictionary containing generated CSV paths.

            The dictionary always contains:

            - `resumen_producto`
            - `resumen_categoria`
            - `resumen_mensual`
            - `resumen_mejores_vendidos_por_mes`
            - `resumen_categoria_mayor_ingreso_por_mes`

            It may also contain:

            - `ciudad_resumen`
            - `metodo_de_pago_resumen`
        """
    reports = {}
    reports["resumen_producto"] = create_save_analysis_result_csv_files_and_path(analysis_result["product_summary"], output_folder, file_name, "productos")
    reports["resumen_categoria"] = create_save_analysis_result_csv_files_and_path(analysis_result["category_summary"], output_folder, file_name, "categorias")
    reports["resumen_mensual"] = create_save_analysis_result_csv_files_and_path(analysis_result["monthly_summary"], output_folder, file_name, "meses")
    reports["resumen_mejores_vendidos_por_mes"] = create_save_analysis_result_csv_files_and_path(analysis_result["monthly_best_selling_product"], output_folder, file_name, "producto_top_mensual")
    reports["resumen_categoria_mayor_ingreso_por_mes"] = create_save_analysis_result_csv_files_and_path(analysis_result["monthly_highest_income_category"], output_folder, file_name, "categoria_top_ingreso_mensual")
    if "city_summary" in analysis_result and analysis_result["city_summary"] is not None:
        reports["ciudad_resumen"] = create_save_analysis_result_csv_files_and_path(analysis_result["city_summary"], output_folder, file_name, "ciudades")
    if "payment_method_summary" in analysis_result and analysis_result["payment_method_summary"] is not None:
        reports["metodo_de_pago_resumen"] = create_save_analysis_result_csv_files_and_path(analysis_result["payment_method_summary"], output_folder, file_name, "metodos_pago")
    return reports

def create_save_analysis_result_csv_files_and_path(
        df: pd.DataFrame, 
        folder: str | Path, 
        file_name: str, 
        prefix: str
) -> Path:
    """Create and save an individual analysis summary as a CSV file.

    Builds the output filename using the shared report base filename and a
    descriptive suffix, creates the destination directory when necessary,
    and exports the provided pandas DataFrame without its index.

    Args:
        df: DataFrame containing the analysis summary to export.
        folder: Directory where the CSV file will be stored.
        file_name: Shared base filename without a file extension.
        prefix: Descriptive filename suffix identifying the analysis summary.

    Returns:
        Path to the generated CSV file.

    Raises:
        ReportSaveError: If the destination directory cannot be created or the
            CSV file cannot be written because of a file-system error.
    """
    output_filename = f"{file_name}_{prefix}.csv"
    try:
        folder = Path(folder)
        folder.mkdir(parents=True, exist_ok=True)
        path = folder / output_filename
        df.to_csv(path, index=False)
    except OSError as error:
        raise ReportSaveError() from error
    return path

def build_sheet_general_summary(wb: Workbook, analysis_result: Dict[str, Any]) -> Worksheet:
    """Build the general sales summary worksheet.

    Creates a worksheet named `Resumen General` and populates it with the main
    sales-processing metrics contained in the analysis result.

    The worksheet contains the total number of rows, valid rows, invalid rows,
    total income, and total units sold.

    Args:
        wb: Excel workbook where the worksheet will be created.
        analysis_result: Dictionary containing the calculated sales metrics.

    Returns:
        The created `Worksheet` containing the general sales summary.
    """
    ws = wb.create_sheet("Resumen General")
    headers = ["Métrica", "Valor"]
    ws.append(headers)
    ws.append(["Total de filas", analysis_result["total_rows"]])
    ws.append(["Filas válidas", analysis_result["total_valid_rows"]])
    ws.append(["Filas inválidas", analysis_result["total_invalid_rows"]])
    ws.append(["Ingreso total", analysis_result["total_income"]])
    ws.append(["Unidades vendidas", analysis_result["total_units_sold"]])
    return ws

def build_sheet(wb: Workbook, df: pd.DataFrame, title_str: str) -> Worksheet:
    """Build a worksheet from a pandas DataFrame.

    Creates a new worksheet using the provided title and writes the complete
    DataFrame into it.

    DataFrame column names are included as worksheet headers, while the
    pandas index is excluded.

    This generic helper is reused for multiple analysis worksheets to avoid
    duplicating DataFrame-to-Excel conversion logic.

    Args:
        wb: Excel workbook where the worksheet will be created.
        df: DataFrame containing the analysis data to export.
        title_str: Title assigned to the new worksheet.

    Returns:
        The created `Worksheet` containing the DataFrame data.
    """
    ws = wb.create_sheet(title_str)
    for row in dataframe_to_rows(df, index=False, header=True):
        ws.append(row)
    return ws

def build_sheet_validation_errors(wb: Workbook, errors: List[Dict[str, Any]]) -> Worksheet:
    """Build the validation-errors worksheet.

    Creates a worksheet named `Validación de errores` containing the
    validation errors detected while processing the source sales data.

    Internal error dictionary keys are mapped to Spanish worksheet headers.
    Each error is written using the predefined column order, and missing
    values are replaced with an empty string.

    Args:
        wb: Excel workbook where the worksheet will be created.
        errors: List of dictionaries containing validation error records.

    Returns:
        The created `Worksheet` containing the validation error records.
    """
    ws = wb.create_sheet("Validación de errores")
    map_headers = {
        "line_number": "línea",
        "column": "columna",
        "error_type": "tipo_error",
        "message": "mensaje",
        "original_value": "valor_original"
    }
    ws.append(list(map_headers.values()))
    for error in errors:
        row = [error.get(value, "") for value in map_headers.keys()]
        ws.append(row)
    return ws

def build_sheet_warnings(wb: Workbook, warnings: List[Dict[str, Any]]) -> Worksheet:
    """Build the validation-warnings worksheet.

    Creates a worksheet named `Advertencias` containing warnings detected
    during sales-data validation.

    Internal warning dictionary keys are mapped to Spanish worksheet headers.
    Each warning is written using the predefined column order, and missing
    values are replaced with an empty string.

    Warning values stored as lists are converted into comma-separated text
    before being written to the worksheet.

    Args:
        wb: Excel workbook where the worksheet will be created.
        warnings: List of dictionaries containing validation warning records.

    Returns:
        The created `Worksheet` containing the validation warning records.
    """
    ws = wb.create_sheet("Advertencias")
    map_headers = {
        "warning_type": "tipo_advertencia",
        "field": "campo",
        "message": "mensaje",
        "affected_value": "valor_afectado",
        "details": "detalles"
    }
    ws.append(list(map_headers.values()))
    for warning in warnings:
        processed_row = []
        for internal_key in map_headers.keys():
            value = warning.get(internal_key, "")
            if isinstance(value, list):
                value = ", ".join(str(v) for v in value)
            processed_row.append(value)
        ws.append(processed_row)
    return ws

def save_report_xlsx(
        analysis_result: Dict[str, Any], 
        errors, 
        warnings, 
        output_folder: str | Path, 
        file_name: str
) -> Path:
    """Generate and save the complete sales analysis as an Excel workbook.

    Creates an XLSX workbook containing general sales metrics, aggregated
    summaries, product rankings, monthly analysis, validation errors, and
    validation warnings.

    The default worksheet created by openpyxl is removed before application
    worksheets are generated.

    The workbook always includes:

    - `Resumen General`
    - `Productos`
    - `Categorías`
    - `Resumen por mes`
    - `Productos mejor vendidos`
    - `Productos con mejor Ingreso`
    - `Producto más vendido por mes`
    - `Categoría mayor ingreso por mes`
    - `Validación de errores`
    - `Advertencias`

    The following worksheets are generated when their corresponding optional
    analysis results are available:

    - `Resumen por ciudad`
    - `Resumen por método de pago`

    Most DataFrame-based worksheets are created through the reusable
    `build_sheet()` helper.

    The function adds the `.xlsx` extension to the shared base filename, creates
    the destination directory when necessary, and saves the completed workbook.

    Args:
        analysis_result: Dictionary containing the complete sales-analysis result
            and summary structures.
        errors: Collection containing validation error records.
        warnings: Collection containing validation warning records.
        output_folder: Directory where the XLSX workbook will be stored.
        file_name: Shared base filename without a file extension.

    Returns:
        Path to the generated XLSX workbook.

    Raises:
        ReportSaveError: If the destination directory cannot be created or the
            workbook cannot be saved because of a file-system error.
    """ 
    output_filename = f"{file_name}.xlsx"
    try:
        folder = Path(output_folder)
        folder.mkdir(parents=True, exist_ok=True)
        path = folder / output_filename

        wb = Workbook()
        wb.remove(wb.active)

        build_sheet_general_summary(wb, analysis_result)

        build_sheet(wb, analysis_result["product_summary"], "Productos")
        build_sheet(wb, analysis_result["category_summary"], "Categorías")
        build_sheet(wb, analysis_result["monthly_summary"], "Resumen por mes")
        if "city_summary" in analysis_result and analysis_result["city_summary"] is not None:
            build_sheet(wb, analysis_result["city_summary"], "Resumen por ciudad")
        if "payment_method_summary" in analysis_result and analysis_result["payment_method_summary"] is not None:
            build_sheet(wb, analysis_result["payment_method_summary"], "Resumen por método de pago")
        df_top_best_products = pd.DataFrame(analysis_result["top_5_best_selling_products"])
        df_top_highest_income_products = pd.DataFrame(analysis_result["top_5_highest_income_products"])
        build_sheet(wb, df_top_best_products, "Productos mejor vendidos")
        build_sheet(wb, df_top_highest_income_products, "Productos con mejor Ingreso")
        build_sheet(wb, analysis_result["monthly_best_selling_product"], "Producto más vendido por mes")
        build_sheet(wb, analysis_result["monthly_highest_income_category"], "Categoría mayor ingreso por mes")
        build_sheet_validation_errors(wb, errors)
        build_sheet_warnings(wb, warnings)
        wb.save(path)
    except OSError as error:
        raise ReportSaveError() from error
    return path
